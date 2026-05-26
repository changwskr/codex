from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("NH_NSIGHT_아키텍처_정의서.docx")


ACCENT = "1F4E79"
ACCENT_DARK = "14395B"
ACCENT_LIGHT = "EAF2F8"
GRAY = "D9E2EC"
TEXT = RGBColor(30, 41, 59)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2EC", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, width_cm=17.0):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(width_cm * 567)))


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 567)))


def apply_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 22, ACCENT_DARK),
        ("Heading 1", 15, ACCENT_DARK),
        ("Heading 2", 12.5, ACCENT),
        ("Heading 3", 10.5, ACCENT_DARK),
    ):
        style = styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(4)


def add_header_footer(section):
    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "NH NSIGHT 통합 아키텍처 정의서"
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.runs[0]
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = "Architecture & Platform Team"
    r = p.runs[0]
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 116, 139)


def add_para(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, ACCENT_LIGHT)
    set_cell_border(cell, "B7CDE0")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    p.add_run("  " + body)
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    if widths:
        set_col_widths(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = header
        set_cell_shading(cell, ACCENT_DARK)
        set_cell_border(cell, ACCENT_DARK)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()
    return table


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_header_footer(section)
    apply_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("NH NSIGHT 통합 아키텍처 정의서")
    sub = doc.add_paragraph()
    sub.add_run("Runtime에서 증명되는 데이터 중심 정보계 아키텍처").bold = True
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)

    add_table(
        doc,
        ["항목", "내용"],
        [
            ["문서 목적", "농협 상호금융 정보계 NSIGHT 아키텍처의 목표, 원칙, 구조, 실행 기준, 검증 기준을 단일 기준선으로 정의한다."],
            ["적용 범위", "마케팅 플랫폼, Single View, RDW, ADW, BI 포털, 배치/ETL, Kafka/CDC, 연계/인터페이스, 운영/보안/DR"],
            ["설계 관점", "Runtime → Mechanism → Physical → Logical → Big Picture로 역추적하고, Big Picture → Runtime으로 구현 검증한다."],
            ["문서 성격", "설계 의사결정 기준서, ARB 심의 기준서, 선도개발 및 본 개발 검증 기준서"],
        ],
        [3.4, 13.6],
    )

    add_note(
        doc,
        "핵심 선언",
        "NSIGHT 아키텍처는 DW 재구축이 아니라, 데이터가 경영 판단과 고객 접점 실행으로 흐르도록 책임·경계·표준·검증 방식을 정렬하는 실행 구조이다.",
    )

    doc.add_heading("1. 아키텍처 비전과 목표", level=1)
    add_para(doc, "NSIGHT 통합 아키텍처의 목표는 기존 정보계의 데이터 저장 중심 구조를 데이터 기반 의사결정 플랫폼으로 전환하는 것이다.")
    add_para(doc, "이를 위해 RDW는 현행성·정합성·Single View 조회를 담당하고, ADW는 분석·집계·경영지표·BI 조회를 담당한다. 마케팅 플랫폼은 고객 접점 실행을 담당하며, BI 포털은 분석 활용과 Self-BI를 담당한다.")
    add_table(
        doc,
        ["전략 축", "정의", "구현 방향"],
        [
            ["Data-Centric", "데이터 성격에 따라 저장·조회·분석 책임을 분리한다.", "RDW/ADW 분리, FAST/DEEP 흐름 분리, 데이터 품질 Gate"],
            ["Scalable", "서비스 증감에 따라 독립 확장 가능한 구조를 만든다.", "AP Scale-Out, Kafka Partition 확장, ETL 병렬화, VM 단위 수평 확장"],
            ["Resilient", "장애가 발생해도 영향 범위를 제한하고 복구 가능한 구조를 만든다.", "AP Active-Active, DB Active-Standby, 센터별 장애 격리, 재로그인 원칙"],
            ["Observable", "거래·성능·장애·데이터 흐름을 측정 가능한 상태로 운영한다.", "GUID, 표준 로그, APM, DB Pool, Kafka Lag, Batch 모니터링"],
        ],
        [3.2, 5.2, 8.6],
    )

    doc.add_heading("2. 아키텍처 설명 방식", level=1)
    add_para(doc, "NSIGHT는 목표에서 그림을 그리는 방식만으로 설명하지 않는다. 실제 운영에서 돌아갈 Runtime 서비스를 먼저 정의하고, 그 서비스를 가능하게 하는 구조를 역방향으로 추적한다.")
    add_table(
        doc,
        ["단계", "관점", "핵심 질문", "주요 산출물"],
        [
            ["1", "Runtime Service", "실제로 어떤 서비스가 운영되는가?", "Runtime 서비스 목록, End-to-End 시나리오"],
            ["2", "Mechanism", "서비스가 안정적으로 동작하려면 어떤 실행 규칙이 필요한가?", "표준 인터페이스, 공통 Header/GUID, 오류/로그/권한 규칙"],
            ["3", "Physical", "실행 규칙을 어떤 시스템에 담을 것인가?", "서버 배치도, Exadata/Cloud/Kafka/CDC/ETL 배치"],
            ["4", "Logical", "물리 구조는 어떤 책임 구조를 따른 것인가?", "레이어 구조, 시스템 책임 정의서, 데이터 소유권 매트릭스"],
            ["5", "Big Picture", "전체 공간과 경계를 어떻게 나눌 것인가?", "통합 아키텍처 조감도, Zone 정의"],
            ["6", "NFR", "품질 기준을 어떻게 보장하는가?", "성능·가용성·확장성·보안성·측정성 검증표"],
        ],
        [1.4, 3.6, 5.2, 6.8],
    )

    doc.add_heading("3. Big Picture와 책임 경계", level=1)
    add_para(doc, "Big Picture의 가치는 그림의 정교함보다 모든 이해관계자가 동의한 경계에 있다. NSIGHT는 사용자/채널, 인터페이스, 데이터 플랫폼, 서비스, 운영/보안/DR 영역을 분리한다.")
    add_table(
        doc,
        ["Zone", "책임", "대표 구성요소"],
        [
            ["Channel / User", "업무 사용자 접속, 전용 단말, 채널 이벤트 발생", "WebTopSuite, 업무 채널, 사용자 브라우저"],
            ["Interface Control", "시스템 간 호출·파일·이벤트·변경데이터 표준 경로 통제", "CruzAPIM, FOS/MFT, API/EAI, CDC, Kafka"],
            ["Data Platform", "현행 데이터와 분석 데이터의 수집·통합·저장·품질 관리", "RDW, ADW, DataStage ETL, 데이터 품질 Gate"],
            ["Service Platform", "고객 접점 실행, Single View, 마케팅 룰, BI 활용", "마케팅 플랫폼, Single View AP, BI Portal"],
            ["Operations / Security / DR", "운영 관측, 보안 통제, 장애 전환, 감사 대응", "APM, 로그, 권한/마스킹, GSLB/L4, DR 절차"],
        ],
        [3.6, 7.2, 6.2],
    )

    doc.add_heading("4. Runtime 기준 아키텍처", level=1)
    add_para(doc, "Runtime은 아키텍처가 실제로 증명되는 지점이다. NSIGHT의 대표 Runtime은 온라인 거래, 배치, 이벤트, 실시간, 대량배치, 분석 조회로 구분한다.")
    add_table(
        doc,
        ["Runtime", "대표 시나리오", "핵심 경로", "검증 기준"],
        [
            ["온라인", "Customer Single View 조회", "UI → CruzAPIM → Online AP → Header/GUID → 권한 → RDW → UI", "p95 3초 이하, 권한/마스킹/로그 검증"],
            ["이벤트", "고객 행동 이벤트 기반 오퍼링", "채널 이벤트 → Kafka → 마케팅 룰 엔진 → 채널 응답", "이벤트 처리 1초 이내, Kafka Lag 모니터링"],
            ["실시간 데이터", "운영계 변경 데이터 반영", "원천 DB → CDC → RDW", "CDC 반영 30초 이내, 건수 대사"],
            ["배치/ETL", "RDW 데이터 정제 후 ADW 적재", "RDW → DataStage ETL → ADW", "배치 윈도우 준수, 재처리 가능성"],
            ["분석/BI", "경영지표·Self-BI 조회", "BI Portal → ADW", "ADW 전용 조회, RDW 영향 없음"],
        ],
        [2.4, 4.2, 6.5, 3.9],
    )

    doc.add_heading("5. 데이터 아키텍처", level=1)
    add_para(doc, "RDW와 ADW의 분리는 DB 분리가 아니라 데이터 성격의 분리이다. 빠르고 정확해야 하는 데이터는 RDW에, 깊고 넓게 분석해야 하는 데이터는 ADW에 둔다.")
    add_table(
        doc,
        ["구분", "RDW", "ADW"],
        [
            ["역할", "현행성·정합성·Single View 조회", "분석·집계·통계·BI 조회"],
            ["데이터", "고객 기본정보, 상품 요약, 최근 거래, 접촉 이력, 실시간 반영 데이터", "장기 이력, 세그먼트, 성과 분석, 분석 마트, 경영지표"],
            ["주요 사용자", "온라인 AP, 마케팅 플랫폼, Single View", "BI Portal, 분석가, 경영지표 서비스"],
            ["금지 원칙", "대용량 분석 쿼리 실행 금지", "운영성 실시간 조회의 직접 의존 금지"],
            ["검증", "Single View 응답시간, RDW SQL 성능, CDC 대사", "BI 조회성능, 집계 정합성, ETL 적재 검증"],
        ],
        [3.0, 7.0, 7.0],
    )
    add_table(
        doc,
        ["흐름", "목적", "처리 경로", "SLA/관리 기준"],
        [
            ["FAST", "고객 행동에 즉시 반응하는 실시간 실행", "채널 이벤트 → Kafka → 마케팅 룰 엔진 → 오퍼링 → RDW 피드백", "이벤트 1초 이내, Kafka Lag 관리"],
            ["DEEP", "정합성 있는 분석과 경영 판단 지원", "원천 → CDC → RDW → ETL/DataStage → ADW → BI Portal", "CDC 30초 이내, 배치 06:00 이전 완료 기준"],
        ],
        [2.4, 4.3, 7.4, 2.9],
    )

    doc.add_heading("6. 인터페이스 및 연계 아키텍처", level=1)
    add_para(doc, "통합은 모든 것을 하나로 합치는 것이 아니라, 각 영역이 자기 책임을 알고 표준 경로로 연결되는 것이다. 직접 DB 접근, 임의 파일 전달, 비표준 호출은 원칙적으로 금지한다.")
    add_table(
        doc,
        ["연계 유형", "적용 대상", "표준 기술/경로", "통제 기준"],
        [
            ["API/EAI", "온라인 조회·업무 서비스 호출", "CruzAPIM, 표준 전문, 공통 Header/GUID", "계약 정의, 인증/권한, 오류코드, 응답시간"],
            ["MFT/File", "대외기관 파일, 대량 파일 전달", "FOS/MFT", "파일 명명, 암호화, 재처리, 수신확인"],
            ["CDC", "운영계 변경 데이터 수집", "CDC 중계서버 → RDW", "반영 지연, 누락 대사, 재처리 기준"],
            ["Kafka", "이벤트 스트리밍", "Topic, Partition, Consumer Group", "Lag, 재처리, 순서성, 보관기간"],
            ["ETL", "정제·변환·적재", "DataStage ETL", "품질 Gate, Checkpoint, 배치 윈도우"],
        ],
        [2.8, 4.4, 5.0, 4.8],
    )

    doc.add_heading("7. 물리 및 DR 아키텍처", level=1)
    add_para(doc, "물리 아키텍처의 핵심은 데이터는 깊게, 서비스는 넓게 배치하는 것이다. 데이터 플랫폼은 Exadata 기반 성능과 정합성을 확보하고, 서비스 플랫폼은 Cloud VM 기반 Scale-Out으로 유연성을 확보한다.")
    add_table(
        doc,
        ["영역", "물리 구성", "설계 의도"],
        [
            ["AP / Service", "농협 프라이빗 클라우드 VM, 센터별 Active-Active", "서비스 수평 확장, 장애 영향 격리, 배포 유연성"],
            ["DB / Data", "Oracle Exadata 기반 RDW/ADW 자원 분리", "분석 쿼리가 온라인 조회를 침해하지 않도록 CPU/Memory/I/O 격리"],
            ["DR", "AP Active-Active, DB Active-Standby", "운영 가능한 현실적 고가용성, 장애 시 복구 절차 단순화"],
            ["Traffic", "GSLB + 센터 L4 + L4 Sticky", "최초 센터 선택, 센터 내부 부하분산, 세션 접근 안정화"],
        ],
        [3.2, 6.4, 7.4],
    )

    doc.add_heading("8. 세션 아키텍처", level=1)
    add_para(doc, "NSIGHT 세션 구조는 GSLB가 최초 센터를 선택하고, WebTopSuite가 선택된 센터 L4에 직접 접근하며, 센터 내부 Tomcat Cluster에서 DeltaManager로 세션을 복제하는 방식이다.")
    add_table(
        doc,
        ["항목", "정의", "설계 기준"],
        [
            ["센터 선택", "GSLB DNS Lookup으로 센터 L4 VIP 반환", "GSLB는 매 요청 Gateway가 아니라 최초 센터 선택 역할"],
            ["센터 유지", "WebTopSuite가 CENTER_ID/CENTER_L4_URL 저장", "매 요청마다 GSLB Lookup하지 않음"],
            ["AP 장애", "동일 센터 내부 AP Cluster에서 DeltaManager 복제 세션 사용", "L4 Sticky와 DeltaManager 병행"],
            ["센터 장애", "타 센터 L4로 전환 후 재로그인 또는 재인증", "센터 간 세션 복제 기본 미적용"],
            ["세션 저장 가능", "userId, branchId, role, authLevel, maskingLevel", "사용자와 권한 판단에 필요한 최소 정보"],
            ["세션 저장 금지", "고객조회 결과, Single View 결과, 거래목록, RDW ResultSet", "메모리·보안·복제부하 방지"],
        ],
        [3.0, 5.6, 8.4],
    )

    doc.add_heading("9. 용량산정 기준", level=1)
    add_para(doc, "용량산정은 전체 세션 수와 동시 요청 TPS를 분리해서 판단한다. 세션 수는 전체 사용자 기반으로 산정하고, TPS·Thread·DB Pool·서버 수는 동시 요청자 기반으로 산정한다.")
    add_table(
        doc,
        ["항목", "기준값", "설명"],
        [
            ["전체 사용자", "21,600명", "3,600개 지점 × 6명"],
            ["설계 세션", "26,000~28,000", "20~30% 여유 반영"],
            ["VM 기준", "8 vCPU / 32GB", "온라인 AP Scale-Out 단위"],
            ["VM당 처리량", "250 TPS", "선도개발 성능테스트로 보정"],
            ["목표 응답시간", "p95 3초 이하", "평균 1.0~1.2초 수준 관리 필요"],
            ["GC", "G1GC, MaxGCPauseMillis=200ms", "Heap 12~14GB 기준 운영 표준"],
        ],
        [4.0, 4.5, 8.5],
    )
    add_table(
        doc,
        ["시나리오", "동시 요청률", "동시 요청자", "목표 TPS", "권장 AP 구성"],
        [
            ["기본 운영", "5%", "1,080명", "360 TPS", "센터당 3대, 총 6대 또는 운영 최소 총 4대"],
            ["피크/확장", "10%", "2,160명", "720 TPS", "센터당 4대, 총 8대"],
            ["스트레스/강화", "15%", "3,240명", "1,080 TPS", "센터당 6대, 총 12대"],
        ],
        [3.0, 2.8, 3.0, 2.6, 5.6],
    )
    add_note(
        doc,
        "적용 조건",
        "720 TPS 기준 센터당 4대·총 8대 구성은 VM당 250 TPS, 평균 응답시간 1.0~1.2초, RDW SQL 평균 100~300ms, Hikari 사용률 70~80% 이하, Tomcat Busy Thread 70~80% 이하가 성능테스트에서 확인될 때 유효하다.",
    )

    doc.add_heading("10. 비기능 요구사항", level=1)
    add_table(
        doc,
        ["비기능", "아키텍처 반영", "검증 지표"],
        [
            ["성능", "온라인/배치 분리, RDW/ADW 분리, DB Pool 분리, SQL 실행 위치 제한", "TPS, 평균/p95/p99 응답시간, SQL Time, Busy Thread"],
            ["가용성", "AP Active-Active, DB Active-Standby, L4/GSLB, Failover/Failback 절차", "장애 전환 시간, 센터 장애 수용 TPS, 재로그인 정책"],
            ["확장성", "AP Scale-Out, Kafka Partition 확장, ETL 병렬화, VM 단위 증설", "VM당 TPS, Partition Lag, 배치 처리시간"],
            ["보안성", "권한·마스킹·감사로그·Cookie 보안·세션 최소화", "권한 검증, 개인정보 마스킹, 감사로그 누락률"],
            ["측정성", "GUID, 표준 로그, APM, DB/Kafka/Batch 모니터링", "거래 추적률, Kafka Lag, Pool Wait, GC Pause"],
        ],
        [2.6, 8.0, 6.4],
    )

    doc.add_heading("11. 설계 원칙과 Guardrail", level=1)
    add_table(
        doc,
        ["원칙", "선언", "금지 패턴"],
        [
            ["특성 기반 설계", "업무 특성에 따라 처리 방식과 저장 위치를 결정한다.", "업무 성격과 무관한 일괄 기술 적용"],
            ["책임 경계 명확화", "시스템·데이터·인터페이스 책임을 명확히 분리한다.", "소유권 없는 공유 DB, 책임 불명 인터페이스"],
            ["데이터 흐름 분리", "FAST/DEEP, RDW/ADW, 온라인/배치를 분리한다.", "RDW 대용량 분석 쿼리, 온라인/배치 동일 자원 경합"],
            ["표준 인터페이스 강제", "API/MFT/CDC/ETL/Kafka를 목적별 표준 경로로 사용한다.", "직접 DB 접근, 임의 파일 전달, 비표준 호출"],
            ["운영 가능한 현실 선택", "이론적 완벽함보다 운영자가 실제 복구 가능한 구조를 선택한다.", "복잡한 센터 간 세션 복제, 검증 없는 DB A-A"],
        ],
        [3.2, 7.2, 6.6],
    )

    doc.add_heading("12. 검증 및 거버넌스", level=1)
    add_para(doc, "아키텍처의 완성 기준은 문서 승인만이 아니라 Runtime 검증이다. Single View 선도개발은 화면 기능 검증이 아니라 아키텍처 실행 가능성 검증으로 수행한다.")
    add_table(
        doc,
        ["검증 영역", "검증 내용", "필수 증적"],
        [
            ["Single View", "UI부터 RDW 조회, 권한, 마스킹, 로그, 오류, 성능까지 수직 검증", "End-to-End 시나리오, 성능 결과, 로그 추적 결과"],
            ["성능", "720 TPS 피크와 1,080 TPS 스트레스 검증", "TPS, 응답시간, CPU, GC, Busy Thread, Pool 사용률"],
            ["데이터", "CDC/RDW/ADW/ETL 정합성 검증", "건수 대사, 품질 Gate 결과, 재처리 결과"],
            ["DR", "센터 장애, AP 장애, Failback 절차 검증", "장애 시나리오 결과, 운영 Runbook"],
            ["ARB", "원칙 위반과 예외 승인 관리", "예외 승인서, 영향 분석표, 보완 계획"],
        ],
        [3.0, 7.5, 6.5],
    )

    doc.add_heading("13. 한 문장 정의", level=1)
    add_note(
        doc,
        "정의",
        "NH NSIGHT 통합 아키텍처는 RDW/ADW, FAST/DEEP, 온라인/배치, AP/DB DR 구조를 책임과 경계 중심으로 분리하고, Single View Runtime에서 성능·보안·관측성·운영 가능성을 검증하는 데이터 중심 실행 아키텍처이다.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
